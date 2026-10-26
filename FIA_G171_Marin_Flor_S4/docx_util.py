#!/usr/bin/env python3
"""
Utilidades de maquetación para la memoria de la Actividad 2.4.

Concentra el estilo visual: tipografías, colores, tablas, cajas de código,
encabezados, pies de página y figuras con pie.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# Paleta
# ---------------------------------------------------------------------------
NAVY = "0F2A44"
NAVY_2 = "163A5C"
BLUE = "1F6FB2"
BLUE_LIGHT = "D6E8F7"
BLUE_BG = "EAF3FB"
ORANGE = "D9782D"
ORANGE_BG = "FDF1E8"
GOLD = "E9C46A"
GREEN = "2E8B57"
RED = "C0392B"
GRAY = "6B7A8A"
GRAY_LIGHT = "E7EDF3"
BORDER = "D5DEE7"
CODE_BG = "F5F7FA"
WHITE = "FFFFFF"

BODY_FONT = "Calibri"
TITLE_FONT = "Montserrat"
MONO_FONT = "Liberation Mono"

PAGE_WIDTH_CM = 16.6


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# ---------------------------------------------------------------------------
# Runs y párrafos
# ---------------------------------------------------------------------------
def _aplicar_fuente(run, font: str, size: float, color: str,
                    bold: bool = False, italic: bool = False) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)


def espaciado_letras(run, twips: int = 40) -> None:
    rpr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(twips))
    rpr.append(sp)


def texto(doc, txt: str, size: float = 10.5, bold: bool = False,
          italic: bool = False, color: str = NAVY,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after: float = 6,
          space_before: float = 0, font: str = BODY_FONT,
          line_spacing: float = 1.15, keep_with_next: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(txt)
    _aplicar_fuente(run, font, size, color, bold, italic)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line_spacing
    pf.keep_with_next = keep_with_next
    return p


def vineta(doc, txt: str, negrita: str | None = None, size: float = 10.5,
           space_after: float = 3):
    p = doc.add_paragraph(style="List Bullet")
    if negrita:
        r1 = p.add_run(negrita)
        _aplicar_fuente(r1, BODY_FONT, size, NAVY, bold=True)
        r2 = p.add_run(txt)
        _aplicar_fuente(r2, BODY_FONT, size, NAVY)
    else:
        r = p.add_run(txt)
        _aplicar_fuente(r, BODY_FONT, size, NAVY)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.1
    return p


def numerada(doc, txt: str, negrita: str | None = None, size: float = 10.5,
             space_after: float = 3):
    p = doc.add_paragraph(style="List Number")
    if negrita:
        r1 = p.add_run(negrita)
        _aplicar_fuente(r1, BODY_FONT, size, NAVY, bold=True)
        r2 = p.add_run(txt)
        _aplicar_fuente(r2, BODY_FONT, size, NAVY)
    else:
        r = p.add_run(txt)
        _aplicar_fuente(r, BODY_FONT, size, NAVY)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.1
    return p


# ---------------------------------------------------------------------------
# Estilos y sección
# ---------------------------------------------------------------------------
def _borde_estilo(style, color: str, sz: int = 8) -> None:
    ppr = style.element.get_or_add_pPr()
    for el in ppr.findall(qn("w:pBdr")):
        ppr.remove(el)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def configurar_estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(NAVY)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "es-ES")
    rpr.append(lang)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    especificaciones = [
        ("Heading 1", 17, NAVY, 18, 8),
        ("Heading 2", 13, BLUE, 13, 5),
        ("Heading 3", 11.5, NAVY, 10, 4),
    ]
    for nombre, size, color, before, after in especificaciones:
        st = doc.styles[nombre]
        st.font.name = TITLE_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(color)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), TITLE_FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.05

    _borde_estilo(doc.styles["Heading 1"], GOLD, sz=10)

    for nombre in ("List Bullet", "List Number"):
        st = doc.styles[nombre]
        st.font.name = BODY_FONT
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(3)


def _campo(p, instr: str, size: float = 8, color: str = GRAY) -> None:
    run = p.add_run()
    _aplicar_fuente(run, TITLE_FONT, size, color, bold=False)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "1"
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    for el in (fld1, it, fld2, t, fld3):
        run._r.append(el)


def _borde_parrafo(p, edge: str, color: str, sz: int = 6, space: int = 2) -> None:
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)
    ppr.append(pbdr)


def configurar_seccion(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True

    # Encabezado (a partir de la página 2)
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.tab_stops.add_tab_stop(
        Cm(PAGE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT
    )
    r1 = hp.add_run("FUNDAMENTOS DE INTELIGENCIA ARTIFICIAL")
    _aplicar_fuente(r1, TITLE_FONT, 7.5, GRAY, bold=True)
    espaciado_letras(r1, 10)
    hp.add_run("\t")
    r2 = hp.add_run("Actividad 2.4 · BFS y DFS")
    _aplicar_fuente(r2, TITLE_FONT, 7.5, GRAY)
    _borde_parrafo(hp, "bottom", BORDER, sz=4, space=4)

    # Pie de página
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(
        Cm(PAGE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT
    )
    r3 = fp.add_run("Julissa Marin Flor · Grupo G171")
    _aplicar_fuente(r3, TITLE_FONT, 7.5, GRAY)
    fp.add_run("\t")
    r4 = fp.add_run("Página ")
    _aplicar_fuente(r4, TITLE_FONT, 7.5, GRAY)
    _campo(fp, "PAGE")
    r5 = fp.add_run(" de ")
    _aplicar_fuente(r5, TITLE_FONT, 7.5, GRAY)
    _campo(fp, "NUMPAGES")
    _borde_parrafo(fp, "top", BORDER, sz=4, space=4)


# ---------------------------------------------------------------------------
# Tablas
# ---------------------------------------------------------------------------
def sombreado_celda(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for el in tcpr.findall(qn("w:shd")):
        tcpr.remove(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def margenes_celda(cell, top: int = 70, bottom: int = 70,
                    left: int = 110, right: int = 110) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for el in tcpr.findall(qn("w:tcMar")):
        tcpr.remove(el)
    mar = OxmlElement("w:tcMar")
    for nombre, valor in (("top", top), ("bottom", bottom),
                          ("left", left), ("right", right)):
        el = OxmlElement(f"w:{nombre}")
        el.set(qn("w:w"), str(valor))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def bordes_celda(cell, color: str = BLUE, sz: int = 24) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for el in tcpr.findall(qn("w:tcBorders")):
        tcpr.remove(el)
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)
    borders.append(left)
    tcpr.append(borders)


def bordes_tabla(table, color: str = BORDER, sz: int = 4,
                  outside: bool = True, inside_h: bool = True,
                  inside_v: bool = True) -> None:
    tblpr = table._tbl.tblPr
    for el in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(el)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        visible = outside if edge in ("top", "left", "bottom", "right") else (
            inside_h if edge == "insideH" else inside_v
        )
        el.set(qn("w:val"), "single" if visible else "nil")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def anchos_tabla(table, widths_cm) -> None:
    table.autofit = False
    table.allow_autofit = False
    tblpr = table._tbl.tblPr
    for el in tblpr.findall(qn("w:tblLayout")):
        tblpr.remove(el)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for row in table.rows:
        for cell, ancho in zip(row.cells, widths_cm):
            cell.width = Cm(ancho)
    for col, ancho in zip(table.columns, widths_cm):
        col.width = Cm(ancho)


def _texto_celda(cell, txt: str, size: float, color: str,
                 bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(txt))
    _aplicar_fuente(r, BODY_FONT, size, color, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def caption_tabla(doc, texto_caption: str) -> None:
    p = doc.add_paragraph()
    if ". " in texto_caption:
        num, resto = texto_caption.split(". ", 1)
        r1 = p.add_run(num + ". ")
        _aplicar_fuente(r1, BODY_FONT, 9, BLUE, bold=True)
        r2 = p.add_run(resto)
        _aplicar_fuente(r2, BODY_FONT, 9, NAVY, bold=False)
    else:
        r = p.add_run(texto_caption)
        _aplicar_fuente(r, BODY_FONT, 9, NAVY, bold=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True


def tabla_bonita(doc, headers, rows, widths, caption: str | None = None,
                 size: float = 9.2, first_col_bold: bool = True,
                 center_cols=()) -> None:
    if caption:
        caption_tabla(doc, caption)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    bordes_tabla(table, color=BORDER, sz=4)
    anchos_tabla(table, widths)

    # Encabezado
    hdr = table.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)
    for i, titulo in enumerate(headers):
        cell = hdr.cells[i]
        sombreado_celda(cell, NAVY)
        margenes_celda(cell, top=80, bottom=80)
        _texto_celda(cell, titulo, size, WHITE, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER)

    # Filas
    for idx, fila in enumerate(rows):
        cells = table.add_row().cells
        for i, valor in enumerate(fila):
            cell = cells[i]
            if idx % 2 == 1:
                sombreado_celda(cell, "F7F9FC")
            margenes_celda(cell)
            align = (WD_ALIGN_PARAGRAPH.CENTER
                     if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
            bold = first_col_bold and i == 0
            color = NAVY if bold else NAVY_2
            _texto_celda(cell, valor, size, color, bold=bold, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(doc, titulo: str, cuerpo: str, color: str = BLUE,
            fondo: str = BLUE_BG) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(table, outside=False, inside_h=False, inside_v=False)
    anchos_tabla(table, [PAGE_WIDTH_CM])
    cell = table.rows[0].cells[0]
    sombreado_celda(cell, fondo)
    bordes_celda(cell, color=color, sz=24)
    margenes_celda(cell, top=140, bottom=140, left=190, right=150)

    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(titulo)
    _aplicar_fuente(r1, TITLE_FONT, 9.5, color, bold=True)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(cuerpo)
    _aplicar_fuente(r2, BODY_FONT, 9.6, NAVY)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------------------
# Código y figuras
# ---------------------------------------------------------------------------
def bloque_codigo(doc, texto_codigo: str, size: float = 7.3,
                  color_borde: str = BLUE, mantener_junto: bool = True) -> None:
    lineas = texto_codigo.rstrip("\n").split("\n")
    total = len(lineas)
    for i, linea in enumerate(lineas):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.right_indent = Cm(0.05)
        p.paragraph_format.keep_with_next = mantener_junto and i < total - 1
        run = p.add_run(linea if linea.strip() else " ")
        _aplicar_fuente(run, MONO_FONT, size, "1C2833")

        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), CODE_BG)
        ppr.append(shd)

        pbdr = OxmlElement("w:pBdr")
        for edge in ("left", "top" if i == 0 else None,
                     "bottom" if i == total - 1 else None):
            if not edge:
                continue
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "18" if edge == "left" else "4")
            el.set(qn("w:space"), "3")
            el.set(qn("w:color"), color_borde if edge == "left" else BORDER)
            pbdr.append(el)
        ppr.append(pbdr)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figura(doc, ruta: Path, caption: str, width_in: float = 6.3) -> None:
    doc.add_picture(str(ruta), width=Inches(width_in))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True

    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if ". " in caption:
        num, resto = caption.split(". ", 1)
        r1 = cp.add_run(num + ". ")
        _aplicar_fuente(r1, BODY_FONT, 8.8, BLUE, bold=True)
        r2 = cp.add_run(resto)
        _aplicar_fuente(r2, BODY_FONT, 8.8, GRAY, italic=True)
    else:
        r = cp.add_run(caption)
        _aplicar_fuente(r, BODY_FONT, 8.8, GRAY, italic=True)
    cp.paragraph_format.space_after = Pt(10)


def linea_horizontal(doc, color: str = GOLD, sz: int = 12,
                     space_before: int = 0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(0)
    _borde_parrafo(p, "bottom", color, sz=sz, space=1)


def metadatos(doc: Document) -> None:
    cp = doc.core_properties
    cp.title = "Actividad 2.4 — Búsqueda ciega (BFS y DFS)"
    cp.subject = "Fundamentos de Inteligencia Artificial"
    cp.author = "Julissa Marin Flor"
    cp.keywords = "BFS, DFS, laberinto, búsqueda ciega, Python"
    cp.comments = "Memoria de la Actividad 2.4 — Universidad Politécnica de Atlautla"
